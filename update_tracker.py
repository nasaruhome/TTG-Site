from __future__ import annotations

import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("The Python package 'openpyxl' is required to refresh the tracker.")
    print("Run: py -m pip install openpyxl")
    raise


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "TTG worksheet for wir and mir.xlsx"
OUTPUT = ROOT / "index.html"
STATUS_OPTIONS = ["Just Created", "Not Submitted", "ITP Need", "With Consultant", "No Updates", "Process", "Code-A", "Code-B", "with Doc Controller"]


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\ufffd", "-").strip()


def value_at(ws, row, col):
    return clean(ws.cell(row, col).value)


def find_header(ws, required):
    for row in range(1, min(ws.max_row, 8) + 1):
        vals = [clean(c.value).lower() for c in ws[row]]
        if all(any(needle in v for v in vals) for needle in required):
            return row
    return None


def status_choice(value, fallback="Process"):
    v = clean(value).lower()
    if v == "pending":
        return "Not Submitted"
    return next((s for s in STATUS_OPTIONS if s.lower() == v), fallback)


def attachment_index():
    found = {"MIR": [], "WIR": []}
    for kind, folder in (("MIR", ROOT / "TTG MIR"), ("WIR", ROOT / "TTG WIR")):
        if folder.exists():
            for path in sorted(folder.rglob("*.pdf"), key=lambda p: p.name.lower()):
                found[kind].append({"name": path.name, "path": path.relative_to(ROOT).as_posix()})
    return found


def record(sheet, row, kind, number="", description="", mir="", wir="", mir_status="Process", wir_status="Process", status="Process", extra=None, source_status=""):
    key = f"{sheet}:{row}:{number or mir or wir or row}"
    return {
        "key": key,
        "sheet": sheet,
        "row": row,
        "kind": kind,
        "number": number,
        "description": description,
        "mir": mir,
        "wir": wir,
        "mirStatus": mir_status,
        "wirStatus": wir_status,
        "status": status,
        "sourceStatus": source_status,
        "extra": extra or [],
    }


def extract_records(workbook):
    out = []

    # Only the main combined status log is in scope. Keep workbook numbers 1–24;
    # repeated number 18 rows are retained as separate rows exactly as supplied.
    ws = workbook["Sheet1"]
    hdr = find_header(ws, ["sr. no", "material", "mir", "wir"])
    if hdr:
        for r in range(hdr + 1, ws.max_row + 1):
            sr, desc, mir, ms, wir, wst = (value_at(ws, r, c) for c in (1, 2, 11, 12, 13, 14))
            if not any((desc, mir, wir)) or not sr:
                continue
            try:
                if int(float(sr)) > 24:
                    continue
            except ValueError:
                continue
            out.append(record(ws.title, r, "MIR & WIR", sr, desc, mir, wir,
                              status_choice(ms), status_choice(wst, "Process" if not wir else "Process"),
                              status_choice(ms), [("MIR Status (workbook)", ms), ("WIR Status (workbook)", wst)]))
    return out


def attach_matches(records, attachments):
    for rec in records:
        rec["files"] = []
        searches = [("MIR", rec.get("mir", "")), ("WIR", rec.get("wir", ""))]
        for kind, number in searches:
            if not number:
                continue
            token = number.casefold()
            for file in attachments[kind]:
                if token in file["name"].casefold():
                    rec["files"].append({**file, "kind": kind})


def main():
    if "--initialize-missing-wir" in sys.argv:
        wb = load_workbook(WORKBOOK, data_only=False)
        ws = wb["Sheet1"]
        wir_files = [p.name.casefold() for p in (ROOT / "TTG WIR").rglob("*.pdf")]
        changed = []
        for row in range(5, ws.max_row + 1):
            excel_no = ws.cell(row, 1).value
            wir_no = clean(ws.cell(row, 13).value)
            if excel_no is None or int(excel_no) > 24:
                continue
            has_pdf = bool(wir_no and any(wir_no.casefold() in name for name in wir_files))
            if not has_pdf and ws.cell(row, 14).value != "Process":
                ws.cell(row, 14).value = "Process"
                changed.append((int(excel_no), row))
        wb.save(WORKBOOK)
        wb.close()
        print(f"Set missing WIR PDF statuses to Process in {len(changed)} workbook rows.")
        return
    if not WORKBOOK.exists():
        sys.exit(f"Workbook not found: {WORKBOOK.name}")
    wb = load_workbook(WORKBOOK, data_only=True, read_only=False)
    data = {"records": extract_records(wb), "attachments": attachment_index(), "workbook": WORKBOOK.name, "logo": "Logo/TR.jpg"}
    attach_matches(data["records"], data["attachments"])
    json_data = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    template = ROOT / "tracker_template.html"
    if not template.exists():
        sys.exit("tracker_template.html is missing.")
    html = template.read_text(encoding="utf-8").replace("__TRACKER_DATA__", json_data)
    OUTPUT.write_text(html, encoding="utf-8")
    print(f"Updated {OUTPUT.name}: {len(data['records'])} records, "
          f"{len(data['attachments']['MIR'])} MIR PDFs, {len(data['attachments']['WIR'])} WIR PDFs.")


if __name__ == "__main__":
    main()
