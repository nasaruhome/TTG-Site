from __future__ import annotations

import json
import re
import sys
from datetime import datetime
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

import update_tracker


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "TTG worksheet for wir and mir.xlsx"
HOST = "127.0.0.1"
PORT = 8765
VALID_STATUSES = {"Just Created", "Not Submitted", "ITP Need", "With Consultant", "No Updates", "Process", "Code-A", "Code-B", "with Doc Controller"}
KEY_RE = re.compile(r"^Sheet1:(\d+):(\d+)$")


def log(message):
    with (ROOT / "tracker_server.log").open("a", encoding="utf-8") as stream:
        stream.write(f"{datetime.now().isoformat(timespec='seconds')} {message}\n")


class TrackerHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def log_message(self, fmt, *args):
        # Keep the background helper quiet.
        return

    def send_json(self, code, value):
        body = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        if self.path != "/update":
            self.send_json(404, {"error": "Unknown endpoint."})
            return
        origin = self.headers.get("Origin", "")
        if origin and origin not in (f"http://{HOST}:{PORT}", f"http://localhost:{PORT}"):
            self.send_json(403, {"error": "Requests must come from the local tracker page."})
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length < 1 or length > 100_000:
                raise ValueError("The status update is empty or too large.")
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            changes = payload.get("changes")
            if not isinstance(changes, list) or not changes:
                raise ValueError("There are no status changes to save.")
            log("Received status update: " + repr([(c.get("key"), c.get("field"), c.get("status")) for c in changes if isinstance(c, dict)]))

            validated = []
            for change in changes:
                if not isinstance(change, dict):
                    raise ValueError("Invalid status change.")
                match = KEY_RE.fullmatch(str(change.get("key", "")))
                field = change.get("field")
                status = change.get("status")
                if not match or field not in ("mirStatus", "wirStatus") or status not in VALID_STATUSES:
                    raise ValueError("A status change has an invalid record or status.")
                row = int(match.group(1))
                excel_no = int(match.group(2))
                if not 5 <= row <= 31 or not 1 <= excel_no <= 24:
                    raise ValueError("A status change refers to a record outside Excel numbers 1–24.")
                validated.append((row, excel_no, field, status))

            workbook = load_workbook(WORKBOOK, data_only=False)
            sheet = workbook["Sheet1"]
            for row, excel_no, field, status in validated:
                if str(sheet.cell(row, 1).value) != str(excel_no):
                    raise ValueError("The workbook row no longer matches this page. Refresh the tracker and try again.")
                target_col = 12 if field == "mirStatus" else 14
                sheet.cell(row, target_col).value = status
            try:
                workbook.save(WORKBOOK)
            except PermissionError as exc:
                raise RuntimeError("Close the workbook in Excel, then press Update Excel again.") from exc
            finally:
                workbook.close()

            update_tracker.main()
            self.send_json(200, {"saved": len(validated)})
        except (ValueError, KeyError, json.JSONDecodeError) as exc:
            log(f"Rejected status update: {exc}")
            self.send_json(400, {"error": str(exc)})
        except RuntimeError as exc:
            log(f"Workbook update unavailable: {exc}")
            self.send_json(409, {"error": str(exc)})
        except PermissionError:
            log("Workbook update failed because the workbook is locked.")
            self.send_json(409, {"error": "The workbook is open or locked. Close it in Excel and try again."})
        except Exception as exc:
            log(f"Unexpected workbook update error: {exc!r}")
            self.send_json(500, {"error": f"Could not save the workbook: {exc}"})


class SingleTrackerServer(ThreadingHTTPServer):
    # Prevent multiple versions from sharing the same port on Windows.
    allow_reuse_address = False


def main():
    try:
        server = SingleTrackerServer((HOST, PORT), TrackerHandler)
    except OSError as exc:
        # A previous launch may already be serving the tracker.
        log(f"Could not bind to {HOST}:{PORT}: {exc}")
        return
    log(f"Serving tracker at http://{HOST}:{PORT}/")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
        log("Tracker helper stopped.")


if __name__ == "__main__":
    main()
